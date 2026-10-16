import os.path
from openpyxl import load_workbook
from util_tools.logs_util.recordlog import logs


class ExcelDataReader:
    """
    读取Excel文件数据
    """

    def __init__(self, file_path):
        self.file_path = file_path
        abspath = os.path.abspath(self.file_path)

        if not os.path.exists(abspath):
            raise FileNotFoundError(f'文件不存在：{abspath}')

        # read_only：控制工作簿是否以制度模式打开，如果为True，工作簿以只读模式打开，不允许对工作簿进行修改
        self.workbook = load_workbook(self.file_path, read_only=False)

    def read_entire_row(self, sheet_name='Sheet1', row_index=1):
        """
        获取Excel文件一整行的数据
        :param sheet_name: 工作表名称
        :param row_index: 返回哪一行的数据，索引从1开始
        :return:
        """
        sheet = self.workbook[sheet_name]
        rows_data = [cell.value for cell in sheet[row_index]]
        return rows_data

    def read_multiple_rows(self, sheet_name='Sheet1'):
        """
        获取Excel文件多行数据。每一个数据组成一个列表
        :param sheet_name:
        :return: 返回list，如，[['1', '2', '3'], ['4', '5', '6'], ['7', '8', '9']]
        """
        try:
            all_rows_data = []
            sheet = self.workbook[sheet_name]
            max_col = sheet.max_column  # 读取最大列数
            for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=max_col):
                row_data = [cell.value for cell in row]
                all_rows_data.append(row_data)
            return all_rows_data
        except Exception as e:
            logs.error('读取Excel文件数据异常,异常原因为{}'.format(e))

    def read_entire_column(self, sheet_name='Sheet1', column_index=1):
        """
        获取Excel文件一整列的数据
        :param sheet_name: 工作表名称
        :param column_index: 读取哪一列的数据，索引从1开始
        :return:
        """
        try:
            column_data = []
            sheet = self.workbook[sheet_name]
            # min_row：从哪一行开始读取，max_row：从哪一行结束读取，min_col：从哪一列开始读取，max_col：从哪一列结束读取
            for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=column_index, max_col=column_index):
                column_data.append(row[0].value)
            return column_data
        except Exception as e:
            logs.error('读取Excel文件数据异常,异常原因为{}'.format(e))

    def read_cell_value(self, row_index=1, column_index=1, sheet_name='Sheet1'):
        """
        获取Excel文件指定单元格的数据
        :param sheet_name: 工作表名称
        :param row_index: 行索引，索引从1开始
        :param column_index: 列索引，索引从1开始
        :return:
        """
        try:
            sheet = self.workbook[sheet_name]
            return sheet.cell(row=row_index, column=column_index).value
        except Exception as e:
            logs.error('读取Excel文件数据异常,异常原因为{}'.format(e))

    def close(self):
        self.workbook.close()
